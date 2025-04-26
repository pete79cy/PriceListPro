import pandas as pd
from openpyxl import load_workbook
import os

print('Excel file size:', os.path.getsize('uploads/quotation_template(13).xlsx'), 'bytes')

# Read the Excel file with pandas
print("\nReading with pandas:")
try:
    df = pd.read_excel('uploads/quotation_template(13).xlsx', engine='openpyxl')
    print("Column headers:", df.columns.tolist())
    print("Data sample:")
    print(df.head())
except Exception as e:
    print("Error reading with pandas:", str(e))

# Read with openpyxl
print("\nReading with openpyxl:")
try:
    wb = load_workbook('uploads/quotation_template(13).xlsx')
    print("Sheet names:", wb.sheetnames)
    
    sheet = wb.active
    print("Sheet dimensions:", sheet.dimensions)
    
    # Get column headers
    headers = []
    for cell in next(sheet.iter_rows()):
        headers.append(cell.value)
    print("Column headers:", headers)
    
    # Print first few rows
    print("\nFirst 5 rows:")
    rows = list(sheet.iter_rows(values_only=True))
    for i, row in enumerate(rows[1:6], 1):
        if any(cell is not None for cell in row):  # Only print non-empty rows
            print(f"Row {i}:", row)
except Exception as e:
    print("Error reading with openpyxl:", str(e))