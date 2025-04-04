"""
Invoice Excel Parser - Functions for parsing customer invoice Excel files.
Handles specific format with merged cells for:
- Customer name in line 11 (columns J-AE)
- From line 16 onwards:
  - Scientific name in merged columns E-L
  - Description in merged columns M-S
  - Quantity in merged columns T-W
  - Price in merged columns X-Z
  - VAT type in merged columns AA-AD (5% or 19%)
  - Total excluding VAT in column AE
"""

import pandas as pd
import re
import os
import uuid
import traceback
from datetime import datetime
from openpyxl import load_workbook
from app import db
from models import Product, Customer, PriceList
from utils.logger import logger
from utils.product_verification import verify_product_exists

def extract_customer_name_from_invoice(excel_path):
    """
    Extract customer name from line 11, merged columns J-AE
    
    Args:
        excel_path (str): Path to the Excel file
        
    Returns:
        str: Customer name or None if not found
    """
    try:
        # Load the workbook with openpyxl to handle merged cells
        wb = load_workbook(excel_path, data_only=True)
        ws = wb.active
        
        # Line 11, column J (index 9 in 0-based indexing)
        customer_cell = ws.cell(row=11, column=10)  # Column J = 10
        customer_name = customer_cell.value
        
        if customer_name:
            # Clean up the customer name
            customer_name = customer_name.strip()
            logger.info(f"Extracted customer name: {customer_name}")
            return customer_name
        else:
            logger.warning("Could not extract customer name from invoice")
            return None
            
    except Exception as e:
        logger.error(f"Error extracting customer name from invoice: {str(e)}")
        return None

def parse_vat_rate(vat_text):
    """
    Parse VAT rate from text like "19%" or "5%"
    
    Args:
        vat_text (str): Text containing VAT rate
        
    Returns:
        float: VAT rate as float (19.0, 5.0) or default 19.0 if not found
    """
    if not vat_text or not isinstance(vat_text, str):
        return 19.0
    
    # Look for percentage in the text
    match = re.search(r'(\d+)%', vat_text)
    if match:
        try:
            return float(match.group(1))
        except (ValueError, TypeError):
            pass
    
    # If we can't parse it, return default
    return 19.0

def parse_invoice_excel(excel_path, customer_id=None):
    """
    Parse customer invoice Excel file with specific format
    
    Args:
        excel_path (str): Path to the Excel file
        customer_id (int, optional): Customer ID if known. If None, will try to find customer by name.
        
    Returns:
        dict: Extracted products and price data
    """
    logger.info(f"Parsing customer invoice Excel: {excel_path}")
    
    try:
        # First, extract customer name if customer_id is not provided
        customer = None
        if not customer_id:
            customer_name = extract_customer_name_from_invoice(excel_path)
            if customer_name:
                # Try to find customer by name
                customer = Customer.query.filter(Customer.name.ilike(f"%{customer_name}%")).first()
                if customer:
                    customer_id = customer.id
                    logger.info(f"Found customer by name: {customer.name} (ID: {customer_id})")
                else:
                    logger.warning(f"Customer not found by name: {customer_name}")
        else:
            # Get customer by ID
            customer = Customer.query.get(customer_id)
        
        if not customer_id:
            raise ValueError("Customer ID not provided and could not be determined from invoice")
        
        # Load the workbook with openpyxl to handle merged cells
        wb = load_workbook(excel_path, data_only=True)
        ws = wb.active
        
        # Initialize results
        products = []
        
        # Start processing from line 16
        row = 16
        while True:
            # Check if we've reached the end of the data
            scientific_name_cell = ws.cell(row=row, column=5)  # Column E = 5
            if not scientific_name_cell.value:
                # Break if no more products
                break
                
            # Extract data from merged cells
            scientific_name = scientific_name_cell.value
            description_cell = ws.cell(row=row, column=13)  # Column M = 13
            description = description_cell.value if description_cell.value else scientific_name
            
            quantity_cell = ws.cell(row=row, column=20)  # Column T = 20
            quantity = quantity_cell.value
            
            price_cell = ws.cell(row=row, column=24)  # Column X = 24
            price = price_cell.value
            
            vat_cell = ws.cell(row=row, column=27)  # Column AA = 27
            vat_text = vat_cell.value
            vat_rate = parse_vat_rate(vat_text)
            
            total_ex_vat_cell = ws.cell(row=row, column=31)  # Column AE = 31
            total_ex_vat = total_ex_vat_cell.value
            
            # Clean and convert data
            try:
                if quantity and isinstance(quantity, (int, float)):
                    quantity = float(quantity)
                else:
                    # Try to extract number from text
                    if quantity and isinstance(quantity, str):
                        match = re.search(r'(\d+(?:\.\d+)?)', quantity)
                        if match:
                            quantity = float(match.group(1))
                        else:
                            quantity = 1.0
                    else:
                        quantity = 1.0
                        
                if price and isinstance(price, (int, float)):
                    price = float(price)
                else:
                    # Try to extract number from text
                    if price and isinstance(price, str):
                        match = re.search(r'(\d+(?:\.\d+)?)', price)
                        if match:
                            price = float(match.group(1))
                        else:
                            price = 0.0
                    else:
                        price = 0.0
            except (ValueError, TypeError):
                quantity = 1.0
                price = 0.0
            
            # Check if product exists in the database
            product_data = {
                'name': description,
                'scientific_name': scientific_name
            }
            
            product, message, is_new, was_created = verify_product_exists(
                product_data,
                create_if_missing=True  # Create product if not found
            )
            
            # Create product entry
            product_entry = {
                'product_id': product.id if product else None,
                'scientific_name': scientific_name,
                'description': description,
                'quantity': quantity,
                'price': price,
                'vat_rate': vat_rate,
                'total_ex_vat': total_ex_vat
            }
            
            products.append(product_entry)
            
            # Move to next row
            row += 1
        
        return {
            'customer_id': customer_id,
            'customer': customer,
            'products': products,
            'invoice_date': datetime.now().strftime('%Y-%m-%d'),
            'source_file': os.path.basename(excel_path)
        }
        
    except Exception as e:
        logger.error(f"Error parsing customer invoice Excel: {str(e)}")
        logger.error(traceback.format_exc())
        raise

def update_price_list_from_invoice(invoice_data):
    """
    Update customer price lists based on invoice data
    
    Args:
        invoice_data (dict): Parsed invoice data from parse_invoice_excel
        
    Returns:
        dict: Summary of updates (updated_count, new_count, error_count)
    """
    logger.info(f"Updating price lists for customer ID: {invoice_data['customer_id']}")
    
    updated_count = 0
    new_count = 0
    error_count = 0
    
    customer_id = invoice_data['customer_id']
    source_file = invoice_data.get('source_file', 'unknown')
    
    for product_entry in invoice_data['products']:
        try:
            product_id = product_entry['product_id']
            price = product_entry['price']
            
            if not product_id or price <= 0:
                logger.warning(f"Skipping price list update for invalid product or price: {product_entry}")
                continue
                
            # Check if price list entry already exists
            existing_price = PriceList.query.filter_by(
                customer_id=customer_id,
                product_id=product_id
            ).first()
            
            if existing_price:
                # Update existing price
                if existing_price.price != price:
                    # Only update if price has changed
                    existing_price.price = price
                    existing_price.updated_at = datetime.utcnow()
                    existing_price.source_file = source_file
                    db.session.add(existing_price)
                    updated_count += 1
                    logger.info(f"Updated price for product ID {product_id} to {price}")
            else:
                # Create new price list entry
                new_price = PriceList(
                    customer_id=customer_id,
                    product_id=product_id,
                    price=price,
                    effective_date=datetime.now().date(),
                    source_file=source_file
                )
                db.session.add(new_price)
                new_count += 1
                logger.info(f"Added new price for product ID {product_id}: {price}")
                
        except Exception as e:
            logger.error(f"Error updating price list for product: {str(e)}")
            error_count += 1
    
    # Commit changes if any updates were made
    if updated_count > 0 or new_count > 0:
        try:
            db.session.commit()
            logger.info(f"Successfully committed {updated_count + new_count} price list changes")
        except Exception as e:
            db.session.rollback()
            logger.error(f"Error committing price list changes: {str(e)}")
            error_count += updated_count + new_count
            updated_count = 0
            new_count = 0
    
    return {
        'updated_count': updated_count,
        'new_count': new_count,
        'error_count': error_count,
        'customer_id': customer_id
    }