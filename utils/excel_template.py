import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from utils.logger import logger

def create_excel_template(output_path):
    """
    Create an Excel template file for price list uploads.
    
    Args:
        output_path (str): Path where the Excel template will be saved
        
    Returns:
        bool: True if successful, False if error
    """
    try:
        logger.info(f"Creating Excel template at: {output_path}")
        
        # Create a workbook and select the active worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Price List Template"
        
        # Define headers
        headers = ["Category", "Name", "Scientific Name", "Pot", "Selling Price"]
        
        # Set up header style
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Write headers and apply styles
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            # Set column width based on header length
            column_letter = get_column_letter(col_num)
            ws.column_dimensions[column_letter].width = max(15, len(header) + 5)
        
        # Add example data
        examples = [
            ["Trees", "Oak Tree", "Quercus robur", "15L", 49.99],
            ["Flowers", "Red Rose", "Rosa 'Red Hybrid'", "2L", 12.50],
            ["Climbers", "Ivy", "Hedera helix", "3L", 8.75]
        ]
        
        # Write example data
        for row_num, example in enumerate(examples, 2):
            for col_num, value in enumerate(example, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # Add instruction row
        instruction_row = len(examples) + 3
        ws.cell(row=instruction_row, column=1, value="Instructions:")
        ws.cell(row=instruction_row, column=1).font = Font(bold=True)
        
        instructions = [
            "1. Fill in your price list data using the format shown in the examples above.",
            "2. 'Category' and 'Scientific Name' are optional but recommended.",
            "3. 'Name' and 'Selling Price' are required fields.",
            "4. 'Pot' should include the pot size (e.g., '2L', '5L', etc.).",
            "5. Save the file as .xlsx or .xls before uploading."
        ]
        
        for i, instruction in enumerate(instructions):
            ws.cell(row=instruction_row + i + 1, column=1, value=instruction)
            ws.merge_cells(f"A{instruction_row + i + 1}:E{instruction_row + i + 1}")
            ws.cell(row=instruction_row + i + 1, column=1).alignment = Alignment(horizontal="left")
        
        # Save the workbook
        wb.save(output_path)
        logger.info(f"Excel template created successfully at: {output_path}")
        return True
        
    except Exception as e:
        logger.error(f"Error creating Excel template: {str(e)}")
        return False

def create_quotation_template(output_path):
    """
    Create an Excel template file for quotation uploads with the format matching the sample file.
    
    Args:
        output_path (str): Path where the Excel template will be saved
        
    Returns:
        bool: True if successful, False if error
    """
    try:
        logger.info(f"Creating quotation Excel template at: {output_path}")
        
        # Create a workbook and select the active worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Quotation Template"
        
        # Define headers based on the sample file format
        headers = [
            "Category", 
            "Description", 
            "Height", 
            "Unit", 
            "Unit price", 
            "Actual Size", 
            "Cost", 
            "Supplier"
        ]
        
        # Set up header style
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Write headers and apply styles
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            # Set column width based on header length
            column_letter = get_column_letter(col_num)
            ws.column_dimensions[column_letter].width = max(15, len(header) + 5)
        
        # Add example data based on the sample file
        examples = [
            ["Trees", "Cupressus sempervirens 'Totem'", "200cm", 4, 45.0, "2m", 35.0, "Shaelos"],
            ["Trees", "Feijoa sellowiana (Multi-stem)", "120/150cm", 4, 20.0, "stem 1,2m", 15.0, "Arocaria"],
            ["Grasses", "Agapanthus africanus", "30cm", 8, 3.5, "2L", "", ""],
            ["Shrubs & Succulents", "Pittosporum tobira 'Nana'", "30cm", 7, 15.0, "5L", 10.0, "Panayiotou"]
        ]
        
        # Write example data
        for row_num, example in enumerate(examples, 2):
            for col_num, value in enumerate(example, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # Add instruction row
        instruction_row = len(examples) + 3
        ws.cell(row=instruction_row, column=1, value="Instructions:")
        ws.cell(row=instruction_row, column=1).font = Font(bold=True)
        
        instructions = [
            "1. Fill in your quotation data using the format shown in the examples above.",
            "2. Column explanations:",
            "   - Category: Plant category (Trees, Grasses, Shrubs, etc.)",
            "   - Description: Plant name/scientific name or description",
            "   - Height: Plant height (e.g., '30cm', '200/250cm', etc.)",
            "   - Unit: Quantity of items",
            "   - Unit price: Selling price per unit",
            "   - Actual Size: Pot size or actual plant dimensions (e.g., '2L', '5L', etc.)",
            "   - Cost: Cost price (optional)",
            "   - Supplier: Supplier name (optional)",
            "3. Only 'Description' and 'Unit price' are required fields.",
            "4. Save the file as .xlsx or .xls before uploading."
        ]
        
        for i, instruction in enumerate(instructions):
            ws.cell(row=instruction_row + i + 1, column=1, value=instruction)
            ws.merge_cells(f"A{instruction_row + i + 1}:H{instruction_row + i + 1}")
            ws.cell(row=instruction_row + i + 1, column=1).alignment = Alignment(horizontal="left")
        
        # Save the workbook
        wb.save(output_path)
        
        logger.info(f"Quotation Excel template created successfully at: {output_path}")
        return True
        
    except Exception as e:
        logger.error(f"Error creating quotation Excel template: {str(e)}")
        return False

def ensure_template_exists(app_static_folder, template_type='price_list'):
    """
    Ensure that the Excel template exists in the static folder.
    If it doesn't exist, create it.
    
    Args:
        app_static_folder (str): Path to the app's static folder
        template_type (str): Type of template to create ('price_list' or 'quotation')
        
    Returns:
        str: Path to the template file
    """
    templates_folder = os.path.join(app_static_folder, 'templates')
    os.makedirs(templates_folder, exist_ok=True)
    
    if template_type == 'quotation':
        template_path = os.path.join(templates_folder, 'quotation_template.xlsx')
        if not os.path.exists(template_path):
            create_quotation_template(template_path)
    else:  # default to price_list
        template_path = os.path.join(templates_folder, 'price_list_template.xlsx')
        if not os.path.exists(template_path):
            create_excel_template(template_path)
    
    return template_path