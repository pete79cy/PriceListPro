"""
Excel generator for price lists using openpyxl.
This module provides functionality to export price list data as Excel files.
"""

import os
import uuid
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from models import Customer, Product, PriceList
from utils.logger import logger

def generate_price_list_excel(customer_id=None, category=None, upload_folder=None):
    """
    Generate an Excel file containing price list data with optional filtering
    
    Args:
        customer_id (int, optional): Filter by customer ID
        category (str, optional): Filter by product category
        upload_folder (str): Folder where the Excel file will be saved
        
    Returns:
        str: Path to the generated Excel file
    """
    try:
        # Create a new workbook and select the active sheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Price List"
        
        # Define headers
        headers = ["Product Name", "Scientific Name", "Category", "Pot", "Customer", "Price (€)", "Last Updated"]
        
        # Set up header style
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Apply header styles
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            
            # Set column width based on header length
            column_letter = get_column_letter(col_num)
            ws.column_dimensions[column_letter].width = max(15, len(header) + 5)
        
        # Query price list data with joins for related information
        from app import db
        query = db.session.query(
            PriceList, Customer, Product
        ).join(
            Customer, PriceList.customer_id == Customer.id
        ).join(
            Product, PriceList.product_id == Product.id
        )
        
        # Apply filters if provided
        if customer_id:
            query = query.filter(PriceList.customer_id == customer_id)
        if category:
            query = query.filter(Product.category == category)
        
        # Order by customer name, product category and name
        query = query.order_by(Customer.name, Product.category, Product.name)
        
        # Execute query
        results = query.all()
        
        # Add data rows
        row_num = 2  # Start from row 2 (after headers)
        for price_list, customer, product in results:
            # Format date
            updated_date = price_list.updated_at.strftime("%Y-%m-%d") if price_list.updated_at else ""
            
            # Add values to the worksheet
            ws.cell(row=row_num, column=1, value=product.name)
            ws.cell(row=row_num, column=2, value=product.scientific_name)
            ws.cell(row=row_num, column=3, value=product.category)
            ws.cell(row=row_num, column=4, value=product.pot)
            ws.cell(row=row_num, column=5, value=customer.name)
            ws.cell(row=row_num, column=6, value=price_list.price)
            ws.cell(row=row_num, column=7, value=updated_date)
            
            # Set alignment for price column
            ws.cell(row=row_num, column=6).alignment = Alignment(horizontal="right")
            
            row_num += 1
        
        # Add thin borders to all cells
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows(min_row=1, max_row=row_num-1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border
        
        # Create a unique filename
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        if customer_id:
            customer = Customer.query.get(customer_id)
            customer_name = f"{customer.name.replace(' ', '_')}_" if customer else ""
        else:
            customer_name = "All_Customers_"
            
        category_part = f"{category.replace(' ', '_')}_" if category else ""
        filename = f"PriceList_{customer_name}{category_part}{timestamp}.xlsx"
        
        # Ensure the upload folder exists
        if not upload_folder:
            upload_folder = 'uploads'
        os.makedirs(upload_folder, exist_ok=True)
        
        # Save the workbook
        file_path = os.path.join(upload_folder, filename)
        wb.save(file_path)
        
        logger.info(f"Generated price list Excel file: {file_path}")
        return file_path
        
    except Exception as e:
        logger.error(f"Error generating price list Excel: {str(e)}")
        raise